import sys
import os
import fitz
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from modules.document_detector import verify_document

SAMPLES_DIR = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(SAMPLES_DIR, exist_ok=True)

def create_authentic_pdf() -> str:
    path = os.path.join(SAMPLES_DIR, "authentic_health_bulletin.pdf")
    doc = fitz.open()
    page = doc.new_page()
    
    text = (
        "GLOBAL HEALTH CLINICAL EVALUATION\n\n"
        "According to a study published in Nature Medicine and confirmed by researchers at "
        "Oxford University and the World Health Organization, public immunization campaigns "
        "demonstrated an 88% reduction in transmission rates across monitored test centers. "
        "Official statements and statistics confirm positive therapeutic outcomes across all demographics."
    )
    page.insert_text((50, 72), text, fontsize=11, fontname="helv")
    
    doc.set_metadata({
        "title": "Global Health Clinical Evaluation",
        "author": "World Health Organization Research Team",
        "creator": "Adobe InDesign 2026",
        "producer": "Adobe PDF Library 15.0",
        "creationDate": "D:20260510100000Z",
        "modDate": "D:20260510100000Z"
    })
    doc.save(path)
    doc.close()
    return path

def create_manipulated_pdf() -> str:
    path = os.path.join(SAMPLES_DIR, "manipulated_forged_press_release.pdf")
    doc = fitz.open()
    page = doc.new_page()
    
    # Sensationalist text with multiple font styles
    page.insert_text((50, 70), "SHOCKING BOMBSHELL LEAK!", fontsize=16, fontname="times-bold")
    page.insert_text((50, 110), "Secret miracle cure 100% suppressed by big pharma and elites!", fontsize=12, fontname="courier")
    page.insert_text((50, 140), "Doctors are furious! Share this now before the video is banned from TV forever!!!", fontsize=11, fontname="helv")
    
    # Metadata indicating graphic design tool + 3-year timeline mismatch
    doc.set_metadata({
        "title": "Confidential Leak",
        "author": "Anonymous Whistleblower",
        "creator": "Canva Design Platform",
        "producer": "Canva PDF Exporter",
        "creationDate": "D:20230101080000Z",
        "modDate": "D:20260912120000Z"
    })
    doc.save(path)
    doc.close()
    return path

def create_sample_docx() -> str:
    path = os.path.join(SAMPLES_DIR, "briefing_memo.docx")
    doc = docx.Document()
    doc.add_heading("Quarterly Scientific Briefing", level=1)
    doc.add_paragraph(
        "According to reports published by Reuters and confirmed by university researchers, "
        "the renewable energy grid expansion has reached 65% capacity milestone this quarter."
    )
    doc.core_properties.author = "Chief Scientist"
    doc.core_properties.title = "Quarterly Scientific Briefing"
    doc.save(path)
    return path

def run_test(title: str, doc_path: str):
    print("\n" + "=" * 75)
    print(f"TEST CASE: {title}")
    print("=" * 75)
    print(f"DOCUMENT PATH: {doc_path}")
    
    result = verify_document(doc_path)
    
    pred = result["prediction"].upper()
    conf = result["confidence_score"]
    weights = result["weights"]
    indicators = result["indicators"]
    
    print(f">> VERDICT:                  [{pred}]")
    print(f">> CONFIDENCE SCORE:         {conf}%")
    print(f">> TEXT CONTENT SCORE (70%): {weights['model_probability'] * 100:.1f}% misinformation")
    print(f">> METADATA RULES     (30%): {weights['rule_score'] * 100:.1f}% tampering markers")
    print(f">> COMPOSITE SCORE:          {weights['composite_score'] * 100:.1f}% fake")
    
    print("\nMETADATA SUMMARY:")
    meta = indicators.get("document_metadata", {})
    for k, v in meta.items():
        print(f"  * {k}: {v}")
        
    print("\nFORENSIC EXPLANATIONS:")
    for idx, exp in enumerate(result["explanation"], 1):
        print(f"  {idx}. {exp}")
        
    return result

if __name__ == "__main__":
    print("Starting Standalone Document Verification Test Suite...")
    
    auth_pdf = create_authentic_pdf()
    manip_pdf = create_manipulated_pdf()
    sample_docx = create_sample_docx()
    
    res_auth = run_test("AUTHENTIC ACADEMIC / CLINICAL PDF", auth_pdf)
    assert res_auth["prediction"] == "real", f"Expected 'real', got {res_auth['prediction']}"
    
    res_manip = run_test("FORGED SENSATIONALIST PDF WITH CANVA & DATE MISMATCH", manip_pdf)
    assert res_manip["prediction"] == "fake", f"Expected 'fake', got {res_manip['prediction']}"
    
    res_docx = run_test("CREDIBLE DOCX MEMORANDUM", sample_docx)
    assert res_docx["prediction"] == "real", f"Expected 'real', got {res_docx['prediction']}"
    
    print("\n" + "=" * 75)
    print("ALL DOCUMENT VERIFICATION TESTS PASSED SUCCESSFULLY!")
    print("=" * 75)
